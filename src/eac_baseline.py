"""EAC baseline: clean a questionnaire sheet, train the six cascade models, test the flow, write one Excel report.

Follows docs/methodology.md (section numbers in the comments refer to it).

Usage: python src/eac_baseline.py --file PATH_TO_XLSX [--sheet 0] [--out outputs/baseline]
"""

import argparse
import json
import re
import sys
import warnings
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import average_precision_score, confusion_matrix, f1_score, precision_score, recall_score
from sklearn.model_selection import StratifiedKFold, train_test_split
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from xgboost import XGBClassifier

warnings.filterwarnings("ignore")

HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------- fixed settings (agreed 2026-10-15)
SEED = 42
TEST_SIZE = 0.20
N_FOLDS = 5
C_GRID = [0.01, 0.1, 1, 10]
XGB_PARAMS = dict(max_depth=3, n_estimators=200, learning_rate=0.1)
MARGIN = 0.01  # "clearly beats" = at least 1 percentage point (plan section 8)
THRESHOLDS = np.round(np.arange(0.01, 1.0, 0.01), 2).tolist() + [1.01]  # 1.01 = never predict yes
RARE_PAIR_MIN = 5  # label pairs with fewer people than this are grouped for the split
BLANK_DROP_PCT = 95.0
SMALL_GROUP = 30  # fairness groups smaller than this are flagged

# ---------------------------------------------------------------- column roles (section 2.1)
ID_COLS = ["diagnostic_response_id", "Link", "Campaign"]
TARGETS = ["final_eac_protect", "final_eac_esc"]
PROTECTED = [
    "What is your age?",
    "What is your gender?",
    "How would you describe your sexual orientation?",
    "Was there any religion in your home growing up?",
    "How would you categorize the religion?",
    "What is the gender of your oldest child?",
    "What is their gender?",
    "Placeholder - Gender of child 3",
    "What is the age of your oldest child?",
    "How old are they?",
    "Placeholder - Age of child 3",
]
AGE_COL = "What is your age?"
GENDER_COL = "What is your gender?"
PIPELINE_LABELS = ["final_osb", "final_asb"]  # the two other final_* columns used by the existing pipeline
SENSITIVE_WORDS = re.compile(r"\bage\b|\bgender\b|sexual orientation|religio", re.I)
AGE_BINS = [-np.inf, 17, 29, 44, 59, np.inf]
AGE_NAMES = ["under 18", "18-29", "30-44", "45-59", "60+"]

# ---------------------------------------------------------------- labels (section 5)
PROTECT_LABELS = ["Anxiety", "Anxiety-Apathy", "Anger", "Anger-Apathy", "Covert Anger", "Covert Anger-Apathy"]
ESC_LABELS = PROTECT_LABELS + ["Frustration"]
ANGER = {"Anger", "Anger-Apathy", "Covert Anger", "Covert Anger-Apathy"}
COVERT = {"Covert Anger", "Covert Anger-Apathy"}
ANGER_BRANCH = ANGER | {"Frustration"}  # Escalating labels that come from the Anger branch
# label -> (anger, covert, apathy, frustration)
FLAGS = {
    "Anxiety": (0, 0, 0, 0),
    "Anxiety-Apathy": (0, 0, 1, 0),
    "Anger": (1, 0, 0, 0),
    "Anger-Apathy": (1, 0, 1, 0),
    "Covert Anger": (1, 1, 0, 0),
    "Covert Anger-Apathy": (1, 1, 1, 0),
    "Frustration": (0, 0, 0, 1),
}
NAMES_BY_FLAGS = {(a, c, p): lab for lab, (a, c, p, f) in FLAGS.items() if not f}

# ---------------------------------------------------------------- overt / covert questions (section 4.2)
OVERT_QUESTIONS = [
    "Would you describe yourself as direct?",
    "Would you describe yourself as desiring for things to be fair or just?",
    "Do you have a pattern of saying things impulsively that you would later regret and feel the desire to fix things or smooth them over?",
    "Even if you feel like you're trying to help, do people blame you for being overly direct, controlling or quick to anger?",
    "Are you quick to correct or step in verbally in a way that often creates conflict or hurt feelings with others?",
    "Do you ask direct or overt questions?",
    "When a boundary has been crossed do you find yourself quick to point it our or stand up for yourself?",
    "Do you get upset if something or someone attempts to derail your idea, plan or goal?",
    "Do you ideally prefer control over a situation but expect things to go wrong?",
]
COVERT_QUESTIONS = [
    "Do you tend to be a people-pleaser or say what someone needs to hear rather than what is actually on your mind?",
    "Do you tend to keep things to yourself because you don't want to upset others or rock the boat?",
    "Is it challenging for you to voice your wants, needs or concerns?",
    "Do you struggle with expressing boundaries or expectations?",
    "Do you often people-please and not say anything until you can't take it anymore, and then say everything that you've been holding in?",
    "Do you go through prolonged periods of keeping things in and staying quiet followed by needing to stand up for yourself or what feels fair?",
    "Do you find it challenging to trust other people's judgment or ability to follow through on plans or a commitment?",
    "Do you replay fights in your head or imagine what you should have said instead?",
    "If you experience apathy and withdraw, do you tend to isolate and willingly self-sabotage?",
]
SCORE_COLS = ["overt_score", "covert_score", "behavioral_bias_ratio"]

# ---------------------------------------------------------------- the six models (section 5.2 and 8)
# elig / pos are written on the label strings so the count check (section 7, step 4) is independent of the flags
MODELS = [
    dict(name="eac_protect_anger", chain="protect", target="p_anger", boost=True, scores=False, rule=False,
         strength="good", elig=("p", lambda p, e: p.notna()), pos=("p", lambda p, e: p.isin(ANGER))),
    dict(name="eac_protect_covert", chain="protect", target="p_covert", boost=True, scores=True, rule=True,
         strength="good", elig=("p", lambda p, e: p.isin(ANGER)), pos=("p", lambda p, e: p.isin(COVERT))),
    dict(name="eac_protect_apathy", chain="protect", target="p_apathy", boost=False, scores=False, rule=False,
         strength="weak", elig=("p", lambda p, e: p.notna()), pos=("p", lambda p, e: p.str.endswith("Apathy"))),
    dict(name="eac_esc_frustration", chain="esc", target="e_frustration", boost=False, scores=False, rule=False,
         force_model=True,  # user decision 2026-10-15: use the trained model even if it does not beat 'always no' by 1 point
         strength="weak", elig=("p", lambda p, e: ~p.isin(ANGER)), pos=("e", lambda p, e: e.eq("Frustration"))),
    dict(name="eac_esc_covert", chain="esc", target="e_covert", boost=False, scores=False, rule=True,
         strength="very weak", elig=("p", lambda p, e: ~p.isin(ANGER) & e.ne("Frustration")),
         pos=("e", lambda p, e: e.isin(COVERT))),
    dict(name="eac_esc_apathy", chain="esc", target="e_apathy", boost=True, scores=False, rule=False,
         strength="good", elig=("e", lambda p, e: e.ne("Frustration")), pos=("e", lambda p, e: e.str.endswith("Apathy"))),
]  # fmt: skip

_log: list[str] = []


def say(msg: str = "") -> None:
    print(msg, flush=True)


def stop(msg: str) -> None:
    """Stop with a clear message (methodology section 1: do not fix quietly)."""
    say(f"\nSTOPPED: {msg}")
    sys.exit(1)


# ============================================================ cleaning helpers
def norm_col(name) -> str:
    """Section 3, step 3: repair broken apostrophes, straighten curly ones, collapse whitespace."""
    s = str(name).replace("�", "'").replace("’", "'").replace("‘", "'").replace("`", "'")
    return re.sub(r"\s+", " ", s).strip()


def clean_cell(v):
    """Strip strings; turn empty or NaN into NA."""
    if pd.isna(v):
        return pd.NA
    if isinstance(v, str):
        v = v.strip()
        return v if v else pd.NA
    return v


def norm_label(v, allowed: dict):
    """Section 3, step 2: map a raw label to its canonical spelling, or return the cleaned text if unknown."""
    if pd.isna(v):
        return pd.NA
    s = re.sub(r"\s*[-–—]\s*", "-", re.sub(r"\s+", " ", str(v))).strip()
    return allowed.get(s.lower(), s)


def classify(s: pd.Series) -> str:
    """Guess a column type from its non-blank values (same rules as eda_eac.py)."""
    vals = s.dropna().astype(str)
    n = len(vals)
    if n == 0:
        return "all_blank"
    if set(vals.str.lower().unique()) <= {"yes", "no"}:
        return "yes_no"
    if pd.to_numeric(vals, errors="coerce").notna().mean() >= 0.95:
        return "numeric"
    if vals.str.contains(r" \| ", regex=True).mean() >= 0.02:
        return "multiselect"
    nunique = vals.nunique()
    if vals.str.len().median() >= 40 and nunique / n > 0.5:
        return "free_text"
    if nunique <= 30:
        return "category"
    return "short_text"


def assign_roles(df: pd.DataFrame) -> pd.DataFrame:
    """Section 2.1: first rule that matches decides. Every column gets a role, an action and a reason."""
    prot = {norm_col(c) for c in PROTECTED}
    rows = []
    for col in df.columns:
        s = df[col]
        ctype = classify(s)
        pct_blank = round(100 * s.isna().mean(), 1)
        n_unique = int(s.dropna().astype(str).nunique())
        if col in ID_COLS:
            role, action, reason = "identifier", "dropped", "identifier (id, link, campaign)"
        elif col in TARGETS:
            role, action, reason = "target", "target", "used as labels"
        elif col in prot:
            role, action, reason = "protected", "dropped", "protected attribute; kept aside for the fairness check"
        elif col.startswith("final_"):
            if pct_blank >= BLANK_DROP_PCT:
                role, action, reason = "blank", "dropped", f"final_* column, {pct_blank}% blank"
            elif n_unique <= 1:
                role, action, reason = "single value", "dropped", "final_* column, only one distinct answer"
            else:
                role, action, reason = "other label", "kept", "final_* column"
        elif pct_blank >= BLANK_DROP_PCT:
            role, action, reason = "blank", "dropped", f"{pct_blank}% blank (95% or more)"
        elif n_unique <= 1:
            role, action, reason = "single value", "dropped", "only one distinct answer"
        elif ctype in ("free_text", "short_text"):
            role, action, reason = "text", "dropped", f"{ctype.replace('_', ' ')}"
        else:
            role, action, reason = "feature", "kept", ""
        rows.append(dict(column=col, role=role, type=ctype, pct_blank=pct_blank, n_unique=n_unique,
                         action=action, reason=reason))  # fmt: skip
    return pd.DataFrame(rows)


# ============================================================ encoding (section 3, steps 6-9)
class Encoder:
    """Turns the kept question columns into 0/1 columns. Everything is learned from the training rows only."""

    def __init__(self, specs: list[tuple[str, str]]):
        self.specs = specs
        self.plan: list[dict] = []
        self.names: list[str] = []
        self.source: dict[str, str] = {}  # encoded column -> original question column

    @staticmethod
    def _split(x) -> set:
        return {t.strip() for t in re.split(r"\s*\|\s*", str(x)) if t.strip()} if pd.notna(x) else set()

    def fit(self, df: pd.DataFrame) -> "Encoder":
        for col, ctype in self.specs:
            s = df[col]
            na = bool(s.isna().any())
            if ctype == "yes_no":
                step = dict(col=col, kind="yes_no", na=na)
            elif ctype == "numeric":
                num = pd.to_numeric(s, errors="coerce")
                vals = num.dropna().to_numpy(dtype=float)
                edges = np.unique(np.quantile(vals, [0.25, 0.5, 0.75])) if len(vals) else np.array([])
                band = np.searchsorted(edges, vals, side="right")
                bands = [b for b in range(len(edges) + 1) if (band == b).any()]
                step = dict(col=col, kind="numeric", edges=edges, bands=bands, na=bool(num.isna().any()))
            elif ctype == "multiselect":
                levels = sorted(set().union(*[self._split(x) for x in s.dropna()])) if s.notna().any() else []
                step = dict(col=col, kind="multi", levels=levels, na=na)
            else:  # category (and any other short code type)
                step = dict(col=col, kind="cat", levels=sorted(s.dropna().astype(str).unique()), na=na)
            self.plan.append(step)
        self.names = list(self._build(df.iloc[:0]))
        return self

    def _build(self, df: pd.DataFrame) -> dict:
        out: dict[str, np.ndarray] = {}

        def add(name, arr, col):
            out[name] = np.asarray(arr, dtype=np.float32)
            self.source[name] = col

        for st in self.plan:
            col, kind = st["col"], st["kind"]
            s = df[col]
            blank = s.isna().to_numpy()
            if kind == "yes_no":
                add(col, s.astype(str).str.lower().eq("yes").to_numpy(), col)
            elif kind == "numeric":
                num = pd.to_numeric(s, errors="coerce")
                blank = num.isna().to_numpy()
                b = np.searchsorted(st["edges"], num.fillna(0).to_numpy(dtype=float), side="right")
                for k in st["bands"]:
                    add(f"{col} (band {k + 1})", (b == k) & ~blank, col)
            elif kind == "cat":
                v = s.astype(str).to_numpy()
                for lev in st["levels"]:
                    add(f"{col} = {lev}", (v == lev) & ~blank, col)
            else:  # multi
                sets = [self._split(x) for x in s]
                for lev in st["levels"]:
                    add(f"{col} [{lev}]", [lev in z for z in sets], col)
            if st["na"]:
                add(f"{col} (not asked)", blank, col)
        return out

    def transform(self, df: pd.DataFrame) -> pd.DataFrame:
        built = self._build(df)
        return pd.DataFrame({k: built[k] for k in self.names}, index=df.index)


# ============================================================ covert / overt scores (section 4.2)
def covert_scores(df: pd.DataFrame) -> pd.DataFrame:
    """YES counts among the 9 overt / 9 covert questions. Blank or missing questions add 0."""

    def yes_count(qs):
        total = np.zeros(len(df), dtype=float)
        for q in qs:
            if q in df.columns:
                total += df[q].fillna("").astype(str).str.strip().str.upper().eq("YES").to_numpy()
        return total

    overt, covert = yes_count([norm_col(q) for q in OVERT_QUESTIONS]), yes_count([norm_col(q) for q in COVERT_QUESTIONS])
    return pd.DataFrame({"overt_score": overt, "covert_score": covert,
                         "behavioral_bias_ratio": overt / (covert + 1)}, index=df.index)  # fmt: skip


# ============================================================ model helpers
def make_lr(C: float):
    return make_pipeline(StandardScaler(), LogisticRegression(C=C, penalty="l2", class_weight="balanced", max_iter=3000))


def make_xgb(y: np.ndarray):
    pos = max(int(y.sum()), 1)
    return XGBClassifier(**XGB_PARAMS, scale_pos_weight=(len(y) - pos) / pos, random_state=SEED,
                         eval_metric="logloss", n_jobs=1, verbosity=0)  # fmt: skip


def n_folds_for(y: np.ndarray) -> int:
    return int(min(N_FOLDS, np.bincount(y, minlength=2).min()))


def oof_probs(make, X: np.ndarray, y: np.ndarray, k: int) -> np.ndarray:
    """Section 8.1: every training person gets one score from a model that did not see them."""
    oof = np.zeros(len(y))
    for tr, va in StratifiedKFold(k, shuffle=True, random_state=SEED).split(X, y):
        est = make()
        est.fit(X[tr], y[tr])
        oof[va] = est.predict_proba(X[va])[:, 1]
    return oof


def best_threshold(p: np.ndarray, y: np.ndarray) -> tuple[float, float]:
    """Highest accuracy; ties go to the threshold closest to 0.5."""
    accs = [(float(((p >= t) == y).mean()), t) for t in THRESHOLDS]
    top = max(a for a, _ in accs)
    t = min((t for a, t in accs if abs(a - top) < 1e-12), key=lambda t: abs(t - 0.5))
    return t, top


def best_cutoff(score: np.ndarray, y: np.ndarray) -> int:
    """Rule: covert if covert_score >= cut-off. Highest accuracy; ties go to the higher cut-off."""
    best, best_acc = 10, -1.0
    for c in range(10, -1, -1):
        acc = float(((score >= c) == y).mean())
        if acc > best_acc + 1e-12:
            best, best_acc = c, acc
    return best


def rule_oof_acc(score: np.ndarray, y: np.ndarray, k: int) -> float:
    """Cross-validated accuracy of the covert-score rule (cut-off chosen on the other folds each time)."""
    hit = 0
    for tr, va in StratifiedKFold(k, shuffle=True, random_state=SEED).split(score.reshape(-1, 1), y):
        c = best_cutoff(score[tr], y[tr])
        hit += int(((score[va] >= c) == y[va]).sum())
    return hit / len(y)


# ============================================================ Excel output
def write_report(path: Path, sheets: dict) -> None:
    """sheets: name -> list of (title or None, DataFrame, with_index)."""
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415

    head_fill = PatternFill("solid", fgColor="D9E1F2")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, blocks in sheets.items():
            row, heads, titles = 0, [], []
            for title, d, with_index in blocks:
                if title:
                    titles.append((row + 1, title))
                    row += 1
                d.to_excel(xw, sheet_name=name, startrow=row, index=with_index)
                heads.append(row + 1)
                row += len(d) + 2
            ws = xw.sheets[name]
            for r, title in titles:
                ws.cell(r, 1).value = title
                ws.cell(r, 1).font = Font(bold=True, size=12)
            for h in heads:
                for cell in ws[h]:
                    if cell.value is not None:
                        cell.font = Font(bold=True)
                        cell.fill = head_fill
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            title_rows = {r for r, _ in titles}
            for col_cells in ws.columns:
                width = max((len(str(c.value)) for c in col_cells if c.value is not None and c.row not in title_rows), default=8)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 60)
            if len(blocks) == 1:
                ws.freeze_panes = "A2"


def pct(x) -> float:
    return round(100 * float(x), 1)


# ============================================================ main
def main() -> None:
    ap = argparse.ArgumentParser(description="EAC baseline (plan.md)")
    ap.add_argument("--file", required=True, help="input .xlsx file")
    ap.add_argument("--sheet", default="0", help="sheet name or index (default 0)")
    ap.add_argument("--out", default=str(HERE.parent / "outputs" / "baseline"), help="output folder")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    outdir = Path(args.out)
    (outdir / "models").mkdir(parents=True, exist_ok=True)
    warnings_list: list[str] = []
    inform_list: list[str] = []

    # ---------------------------------------------------------------- load and clean (section 3, steps 2-3, 6-7)
    sheet = int(args.sheet) if args.sheet.isdigit() else args.sheet
    raw = pd.read_excel(args.file, sheet_name=sheet, dtype=object)
    fixed = sum(norm_col(c) != re.sub(r"\s+", " ", str(c)).strip() for c in raw.columns)
    raw.columns = [norm_col(c) for c in raw.columns]
    dup = raw.columns[raw.columns.duplicated()].tolist()
    if dup:
        stop(f"duplicate column names after header cleaning: {dup}")
    for t in TARGETS:
        if t not in raw.columns:
            stop(f"target column '{t}' is missing from the file")
    df = raw.map(clean_cell)
    n_src = len(df)
    if fixed:
        warnings_list.append(f"{fixed} column header(s) had broken or curly apostrophes and were repaired.")
    say(f"Loaded {n_src} rows, {df.shape[1]} columns from {Path(args.file).name}")

    allowed_p = {x.lower(): x for x in PROTECT_LABELS}
    allowed_e = {x.lower(): x for x in ESC_LABELS}
    p_all = df["final_eac_protect"].map(lambda v: norm_label(v, allowed_p))
    e_all = df["final_eac_esc"].map(lambda v: norm_label(v, allowed_e))
    bad = []
    for nm, ser, ok in [("final_eac_protect", p_all, PROTECT_LABELS), ("final_eac_esc", e_all, ESC_LABELS)]:
        wrong = ser[ser.isna() | ~ser.isin(ok)]
        if len(wrong):
            bad.append(f"{nm}: {wrong.fillna('(blank)').value_counts().to_dict()}")
    if bad:
        stop("labels outside the allowed values: " + "; ".join(bad))

    # ---------------------------------------------------------------- step 1: exclude rows that break the branch rule
    broke = p_all.isin(ANGER) == e_all.isin(ANGER_BRANCH)
    id_col = "diagnostic_response_id" if "diagnostic_response_id" in df.columns else None
    ids_all = df[id_col].astype(str) if id_col else pd.Series(df.index + 2, index=df.index).astype(str)
    excluded = pd.DataFrame({
        "excel_row": df.index[broke] + 2,
        "id": ids_all[broke].values,
        "protect_label": p_all[broke].values,
        "escalating_label": e_all[broke].values,
        "reason": ["Protect anger-family label with an Escalating Anger-branch label" if pr in ANGER
                   else "Protect not-anger label with an Escalating Anxiety-branch label" for pr in p_all[broke]],
    })  # fmt: skip
    df, p, e, ids = df.loc[~broke], p_all[~broke], e_all[~broke], ids_all[~broke]
    n = len(df)
    say(f"Excluded {int(broke.sum())} rows that break the branch rule; {n} rows used.")

    # ---------------------------------------------------------------- column roles (section 2.1)
    roles = assign_roles(df)
    kept = roles[roles["action"] == "kept"]
    specs = [(r.column, "category" if r.type in ("short_text", "free_text", "all_blank") else r.type)
             for r in kept.itertuples()]  # fmt: skip
    other7 = [c for c in kept[kept["role"] == "other label"]["column"] if c not in PIPELINE_LABELS]
    prot_set = {norm_col(c) for c in PROTECTED}
    for r in roles[~roles["role"].isin(["protected", "identifier", "target"])].itertuples():
        if SENSITIVE_WORDS.search(r.column) and r.column not in prot_set:
            warnings_list.append(f"Review: not on the protected list but mentions age, gender, orientation or religion: '{r.column}' ({r.action})")
    absent_q = [q for q in OVERT_QUESTIONS + COVERT_QUESTIONS if norm_col(q) not in df.columns]
    if absent_q:
        warnings_list.append(f"{len(absent_q)} overt/covert question(s) not found in the data (counted as NO): {absent_q}")
    say(f"Column roles: {roles['role'].value_counts().to_dict()} -> {len(specs)} input columns")

    # ---------------------------------------------------------------- flags (section 5.1)
    flag = pd.DataFrame({
        "p_anger": p.map(lambda x: FLAGS[x][0]), "p_covert": p.map(lambda x: FLAGS[x][1]),
        "p_apathy": p.map(lambda x: FLAGS[x][2]),
        "e_covert": e.map(lambda x: FLAGS[x][1]), "e_apathy": e.map(lambda x: FLAGS[x][2]),
        "e_frustration": e.map(lambda x: FLAGS[x][3]),
    }).astype(int)  # fmt: skip

    # ---------------------------------------------------------------- split (section 6)
    pair = p + " | " + e
    counts = pair.value_counts()
    strata = pair.where(pair.map(counts) >= RARE_PAIR_MIN, "rare")
    if (strata == "rare").sum() == 1:  # one lone rare row cannot be stratified: join the smallest group
        small = counts[counts >= RARE_PAIR_MIN].idxmin()
        strata = strata.where(strata != "rare", small)
    train_idx, test_idx = train_test_split(df.index, test_size=TEST_SIZE, random_state=SEED, stratify=strata)
    train_idx, test_idx = pd.Index(sorted(train_idx)), pd.Index(sorted(test_idx))
    say(f"Split: {len(train_idx)} train / {len(test_idx)} test (seed {SEED})")

    # ---------------------------------------------------------------- encode (section 3, steps 8-9; fit on train only)
    enc = Encoder(specs).fit(df.loc[train_idx])
    Xall = enc.transform(df)
    names = enc.names
    final7_cols = [i for i, nm in enumerate(names) if enc.source[nm] in other7]
    keep_wo = [i for i in range(len(names)) if i not in set(final7_cols)]
    Xnp = Xall.to_numpy(dtype=np.float32)
    pos_of = {ix: i for i, ix in enumerate(df.index)}
    say(f"Encoded: {len(names)} model columns ({len(final7_cols)} from the 7 other final_* columns)")

    # ---------------------------------------------------------------- per-model data sets (section 7)
    sel = {"p": p, "e": e}
    res: list[dict] = []
    bundles: dict[str, dict] = {}
    check_fail = []
    scores_for = lambda idx: covert_scores(df.loc[idx])  # noqa: E731

    for cfg in MODELS:
        src, ef = cfg["elig"]
        psrc, pf = cfg["pos"]
        elig = ef(p, e).fillna(False).astype(bool)
        # label-derived positives, then the flag-derived target must agree (count check)
        pos = (pf(p, e).fillna(False) & elig).astype(bool)
        y_all = flag[cfg["target"]].astype(bool) & elig
        e_tr, e_te = elig.loc[train_idx], elig.loc[test_idx]
        n_tr, n_te = int(e_tr.sum()), int(e_te.sum())
        k_tr, k_te = int(pos.loc[train_idx][e_tr].sum()), int(pos.loc[test_idx][e_te].sum())
        if (n_tr + n_te, k_tr + k_te) != (int(elig.sum()), int(pos.sum())) or not (y_all == pos).all():
            check_fail.append(cfg["name"])
        cfg.update(n_tr=n_tr, n_te=n_te, k_tr=k_tr, k_te=k_te, elig=elig, y=y_all.astype(int))
    if check_fail:
        stop(f"count check failed (section 7, step 4) for: {check_fail}")
    say("Count check passed: eligible rows and positives match the labels for all six models.")

    # ---------------------------------------------------------------- train and choose (section 8)
    for cfg in MODELS:
        name = cfg["name"]
        tr_ix = train_idx[cfg["elig"].loc[train_idx].to_numpy()]
        te_ix = test_idx[cfg["elig"].loc[test_idx].to_numpy()]
        ytr, yte = cfg["y"].loc[tr_ix].to_numpy(), cfg["y"].loc[te_ix].to_numpy()
        rtr = np.array([pos_of[i] for i in tr_ix])
        rte = np.array([pos_of[i] for i in te_ix])
        sc_tr = scores_for(tr_ix) if (cfg["scores"] or cfg["rule"]) else None
        sc_te = scores_for(te_ix) if (cfg["scores"] or cfg["rule"]) else None
        extra_tr = sc_tr.to_numpy() if cfg["scores"] else np.zeros((len(tr_ix), 0))
        Xtr = np.hstack([Xnp[rtr], extra_tr])
        Xtr_wo = np.hstack([Xnp[rtr][:, keep_wo], extra_tr])
        input_cols = names + (SCORE_COLS if cfg["scores"] else [])
        base_cv = float((ytr == 0).mean())  # CV accuracy of "always no"
        row = dict(model=name, strength=cfg["strength"], eligible_train=cfg["n_tr"], positives_train=cfg["k_tr"],
                   eligible_test=cfg["n_te"], positives_test=cfg["k_te"],
                   count_check="ok")  # fmt: skip
        note = []
        k = n_folds_for(ytr)
        bundle = dict(model=None, method="always no", threshold=None, cutoff=None, uses_scores=cfg["scores"],
                      input_columns=input_cols, settings="always no")  # fmt: skip
        if k < 2:
            note.append("fewer than 2 positives or negatives in the training rows: always no")
            inform_list.append(f"{name}: too few positives to train; using 'always no'.")
            cv = dict(lr=np.nan, xgb=np.nan, rule=np.nan, lr_wo=np.nan, chosen=base_cv)
        else:
            # logistic regression: pick C by cross-validation accuracy at the best threshold
            lr_runs = []
            for C in C_GRID:
                pr = oof_probs(lambda C=C: make_lr(C), Xtr, ytr, k)
                t, a = best_threshold(pr, ytr)
                lr_runs.append((a, -C, C, t, pr))
            a_lr, _, C_best, t_lr, pr_lr = max(lr_runs, key=lambda r: (r[0], r[1]))
            pr_wo = oof_probs(lambda: make_lr(C_best), Xtr_wo, ytr, k)
            a_wo = best_threshold(pr_wo, ytr)[1]
            cand = dict(method="logistic regression", acc=a_lr, thr=t_lr, probs=pr_lr, settings=f"L2, balanced, C={C_best}",
                        make=lambda C=C_best: make_lr(C))  # fmt: skip
            a_xgb = np.nan
            if cfg["boost"]:
                pr_x = oof_probs(lambda: make_xgb(ytr), Xtr, ytr, k)
                t_x, a_xgb = best_threshold(pr_x, ytr)
                if a_xgb >= a_lr + MARGIN:
                    cand = dict(method="XGBoost", acc=a_xgb, thr=t_x, probs=pr_x,
                                settings=f"depth {XGB_PARAMS['max_depth']}, {XGB_PARAMS['n_estimators']} trees, lr {XGB_PARAMS['learning_rate']}, balanced",
                                make=lambda: make_xgb(ytr))  # fmt: skip
            a_rule, cutoff = np.nan, None
            if cfg["rule"]:
                s_tr = sc_tr["covert_score"].to_numpy()
                a_rule = rule_oof_acc(s_tr, ytr, k)
                cutoff = best_cutoff(s_tr, ytr)
            usable_model = cand["acc"] >= base_cv + MARGIN
            usable_rule = cfg["rule"] and a_rule >= base_cv + MARGIN
            if cfg["rule"]:
                if usable_model and cand["acc"] >= a_rule + MARGIN:
                    choice = "model"
                elif usable_rule:
                    choice = "rule"
                    note.append("model did not beat the covert-score rule by 1 point: rule used")
                elif usable_model:
                    choice = "model"
                    note.append("model beats 'always no' but the rule does not: model kept")
                else:
                    choice = "none"
            else:
                choice = "model" if usable_model else "none"
            if choice == "none" and cfg.get("force_model"):
                choice = "model"
                msg = f"model used by decision although it does not beat 'always no' by 1 point (CV {pct(cand['acc'])}% vs {pct(base_cv)}%)"
                note.append(msg)
                inform_list.append(f"{name}: {msg}.")
            if choice == "model":
                est = cand["make"]()
                est.fit(Xtr, ytr)
                bundle.update(model=est, method=cand["method"], threshold=cand["thr"], settings=cand["settings"])
                chosen_cv = cand["acc"]
            elif choice == "rule":
                bundle.update(method="covert-score rule", cutoff=cutoff, settings=f"covert if covert_score >= {cutoff}")
                chosen_cv = a_rule
            else:
                chosen_cv = base_cv
                why = f"neither the model (CV {pct(cand['acc'])}%) nor a rule beat 'always no' ({pct(base_cv)}%) by 1 point"
                note.append("always no: " + why)
                inform_list.append(f"{name}: {why}. Using 'always no' (this model never predicts a yes).")
            cv = dict(lr=a_lr, xgb=a_xgb, rule=a_rule, lr_wo=a_wo, chosen=chosen_cv)
        bundles[name] = bundle

        # test numbers on the eligible test rows (true labels pick the rows)
        pred, prob = predict_one(bundle, Xnp[rte], sc_te)
        test_acc = float((pred == yte).mean()) if len(yte) else np.nan
        sec = {}
        if bundle["model"] is not None and yte.sum() > 0:
            sec["test_pr_auc"] = pct(average_precision_score(yte, prob))
        if pred.sum() > 0:
            sec["test_precision"] = pct(precision_score(yte, pred, zero_division=0))
        sec["test_recall"] = pct(recall_score(yte, pred, zero_division=0)) if yte.sum() > 0 else np.nan
        if cfg["k_te"] < 12:
            note.append(f"only {cfg['k_te']} positives in the test rows: trust the cross-validation numbers more")
        row.update(
            method=bundle["method"], settings=bundle["settings"],
            threshold=bundle["threshold"] if bundle["threshold"] is not None else (f"cut-off {bundle['cutoff']}" if bundle["cutoff"] is not None else "n/a"),
            cv_acc_always_no=pct(base_cv), cv_acc_logreg=pct(cv["lr"]) if cv["lr"] == cv["lr"] else np.nan,
            cv_acc_xgboost=pct(cv["xgb"]) if cv["xgb"] == cv["xgb"] else np.nan,
            cv_acc_rule=pct(cv["rule"]) if cv["rule"] == cv["rule"] else np.nan,
            cv_acc_logreg_without_7_final=pct(cv["lr_wo"]) if cv["lr_wo"] == cv["lr_wo"] else np.nan,
            cv_acc_chosen=pct(cv["chosen"]), test_acc_always_no=pct((yte == 0).mean()) if len(yte) else np.nan,
            test_acc_chosen=pct(test_acc) if test_acc == test_acc else np.nan, **sec, note="; ".join(note),
        )  # fmt: skip
        res.append(row)
        say(f"  {name}: {bundle['method']} | CV {pct(cv['chosen'])}% vs always-no {pct(base_cv)}% | test {pct(test_acc)}%")
        joblib.dump(bundle, outdir / "models" / f"{name}.joblib")

    joblib.dump(dict(encoder=enc, columns=names), outdir / "models" / "encoder.joblib")
    (outdir / "models" / "models_summary.json").write_text(json.dumps({
        b: dict(method=v["method"], settings=v["settings"], threshold=v["threshold"], cutoff=v["cutoff"],
                input_columns=v["input_columns"]) for b, v in bundles.items()}, indent=1), encoding="utf-8")  # fmt: skip

    # ---------------------------------------------------------------- test the full flow (section 7 and 9)
    pr_lab, es_lab = run_flow(df, test_idx, Xnp, pos_of, bundles)
    p_te, e_te_true = p.loc[test_idx], e.loc[test_idx]
    maj_p, maj_e = p.loc[train_idx].mode()[0], e.loc[train_idx].mode()[0]
    acc_p, acc_e = float((pr_lab == p_te).mean()), float((es_lab == e_te_true).mean())
    dumb_p, dumb_e = float((p_te == maj_p).mean()), float((e_te_true == maj_e).mean())
    say(f"Flow on test: Protect {pct(acc_p)}% (most common label {pct(dumb_p)}%), Escalating {pct(acc_e)}% ({pct(dumb_e)}%)")
    flow_note = []
    for nm, a, d in [("Protect", acc_p, dumb_p), ("Escalating", acc_e, dumb_e)]:
        if a < d + MARGIN:
            flow_note.append(f"{nm} flow does not clearly beat always predicting the most common label")
            inform_list.append(f"Whole flow ({nm}): accuracy {pct(a)}% is not at least 1 point above the most-common-label accuracy {pct(d)}%.")
    flow_df = pd.DataFrame([
        dict(chain="Protect", flow_accuracy=pct(acc_p), most_common_label=maj_p, most_common_label_accuracy=pct(dumb_p),
             clearly_beats=("yes" if acc_p >= dumb_p + MARGIN else "no"), test_people=len(test_idx),
             macro_f1=pct(f1_score(p_te, pr_lab, average="macro", labels=PROTECT_LABELS, zero_division=0))),
        dict(chain="Escalating", flow_accuracy=pct(acc_e), most_common_label=maj_e, most_common_label_accuracy=pct(dumb_e),
             clearly_beats=("yes" if acc_e >= dumb_e + MARGIN else "no"), test_people=len(test_idx),
             macro_f1=pct(f1_score(e_te_true, es_lab, average="macro", labels=ESC_LABELS, zero_division=0))),
    ])  # fmt: skip

    def label_tables(true, pred, labels):
        ct = pd.DataFrame(confusion_matrix(true, pred, labels=labels), index=labels, columns=labels)
        ct.index.name = "true \\ predicted"
        rec = pd.DataFrame({"label": labels,
                            "people": [int((true == lab).sum()) for lab in labels],
                            "recall_%": [pct((pred[true == lab] == lab).mean()) if (true == lab).any() else np.nan for lab in labels]})  # fmt: skip
        return ct, rec

    ct_p, rec_p = label_tables(p_te, pr_lab, PROTECT_LABELS)
    ct_e, rec_e = label_tables(e_te_true, es_lab, ESC_LABELS)

    # ---------------------------------------------------------------- fairness (section 10)
    fair_blocks = []
    if AGE_COL in df.columns and GENDER_COL in df.columns:
        age = pd.to_numeric(df.loc[test_idx, AGE_COL], errors="coerce")
        age_grp = pd.cut(age, AGE_BINS, labels=AGE_NAMES).astype(object).where(age.notna(), "not given")
        gender = df.loc[test_idx, GENDER_COL].astype(object).where(df.loc[test_idx, GENDER_COL].notna(), "not given")
        for title, grp in [("Accuracy by gender (test people)", gender), ("Accuracy by age group (test people)", age_grp)]:
            rows = []
            for g in sorted(grp.dropna().unique(), key=lambda x: (AGE_NAMES.index(x) if x in AGE_NAMES else 99, str(x))):
                m = (grp == g).to_numpy()
                rows.append(dict(group=g, people=int(m.sum()),
                                 protect_accuracy=pct((pr_lab.to_numpy()[m] == p_te.to_numpy()[m]).mean()),
                                 escalating_accuracy=pct((es_lab.to_numpy()[m] == e_te_true.to_numpy()[m]).mean()),
                                 note="small group: treat with care" if m.sum() < SMALL_GROUP else ""))  # fmt: skip
            fair_blocks.append((title, pd.DataFrame(rows), False))
    else:
        warnings_list.append("Fairness sheet skipped: age or gender column not found.")

    # ---------------------------------------------------------------- report (section 2.2)
    models_df = pd.DataFrame(res)
    test_pred = pd.DataFrame({"id": ids.loc[test_idx].values, "true_protect": p_te.values, "predicted_protect": pr_lab.values,
                              "protect_correct": (p_te.values == pr_lab.values),
                              "true_escalating": e_te_true.values, "predicted_escalating": es_lab.values,
                              "escalating_correct": (e_te_true.values == es_lab.values)})  # fmt: skip
    roles_out = roles.rename(columns={"pct_blank": "pct_blank_%"})
    src_count = pd.Series(enc.source).value_counts()
    roles_out["model_columns"] = roles_out["column"].map(src_count).fillna(0).astype(int)
    summary_kv = pd.DataFrame([
        ("Input file", Path(args.file).name), ("Sheet", str(sheet)),
        ("Rows in file", n_src), ("Rows excluded (branch rule broken)", int(broke.sum())), ("Rows used", n),
        ("Train / test people", f"{len(train_idx)} / {len(test_idx)} (seed {SEED})"),
        ("Columns in file", len(roles)), *[(f"Columns with role '{r}'", int(c)) for r, c in roles["role"].value_counts().items()],
        ("Input columns kept (questions + other final_*)", len(specs)),
        ("Model columns after turning answers into numbers", len(names)),
        ("Flow: Protect accuracy vs most common label (%)", f"{pct(acc_p)} vs {pct(dumb_p)} ({maj_p})"),
        ("Flow: Escalating accuracy vs most common label (%)", f"{pct(acc_e)} vs {pct(dumb_e)} ({maj_e})"),
    ], columns=["item", "value"])  # fmt: skip
    headline = models_df[["model", "method", "cv_acc_always_no", "cv_acc_chosen", "test_acc_always_no", "test_acc_chosen"]]
    inform_df = pd.DataFrame({"inform": inform_list or ["Nothing: every model beat 'always no' by 1+ point and the flow beat the most-common-label baseline."]})
    warn_df = pd.DataFrame({"warning": warnings_list or ["none"]})
    sheets = {
        "Summary": [(None, summary_kv, False), ("Headline accuracy (%)", headline, False),
                    ("Please read: models that fell back or did not clearly beat the baseline", inform_df, False),
                    ("Warnings", warn_df, False)],
        "Columns": [(None, roles_out, False)],
        "Excluded rows": [(None, excluded if len(excluded) else pd.DataFrame({"note": ["No rows excluded"]}), False)],
        "Models": [(None, models_df, False)],
        "Flow": [("Whole flow on the test people", flow_df, False),
                 ("Protect: true (rows) vs predicted (columns)", ct_p, True), ("Protect: recall per label", rec_p, False),
                 ("Escalating: true (rows) vs predicted (columns)", ct_e, True), ("Escalating: recall per label", rec_e, False)],
        "Fairness": fair_blocks or [(None, pd.DataFrame({"note": ["skipped"]}), False)],
        "Test predictions": [(None, test_pred, False)],
    }  # fmt: skip
    write_report(outdir / "eac_baseline_report.xlsx", sheets)
    say(f"\nWrote {outdir / 'eac_baseline_report.xlsx'} and models in {outdir / 'models'}")
    if inform_list:
        say("\nPLEASE NOTE (also on the Summary sheet):")
        for m in inform_list:
            say(f"  - {m}")


def predict_one(bundle: dict, X: np.ndarray, scores: pd.DataFrame | None):
    """Yes/no (0/1) for each row, plus the model's probability when there is one."""
    n = len(X)
    if n == 0:
        return np.zeros(0, dtype=int), None
    if bundle["method"] == "always no":
        return np.zeros(n, dtype=int), None
    if bundle["method"] == "covert-score rule":
        return (scores["covert_score"].to_numpy() >= bundle["cutoff"]).astype(int), None
    Xm = np.hstack([X, scores.to_numpy()]) if bundle["uses_scores"] else X
    prob = bundle["model"].predict_proba(Xm)[:, 1]
    return (prob >= bundle["threshold"]).astype(int), prob


def run_flow(df, idx, Xnp, pos_of, bundles):
    """Section 9: Protect then Escalating, each step using the PREDICTED decision of the step before."""
    rows = np.array([pos_of[i] for i in idx])
    X = Xnp[rows]
    n = len(idx)
    def run(name, mask):
        out = np.zeros(n, dtype=int)
        if mask.any():
            needs = bundles[name]["uses_scores"] or bundles[name]["method"] == "covert-score rule"
            sc = covert_scores(df.loc[idx[mask]]) if needs else None
            out[mask] = predict_one(bundles[name], X[mask], sc)[0]
        return out

    everyone = np.ones(n, dtype=bool)
    anger = run("eac_protect_anger", everyone).astype(bool)
    p_apathy = run("eac_protect_apathy", everyone).astype(bool)
    covert = run("eac_protect_covert", anger).astype(bool)
    protect = []
    for a, c, ap in zip(anger, covert, p_apathy):
        protect.append(NAMES_BY_FLAGS[(int(a), int(a and c), int(ap))])
    # Escalating: the predicted Protect anger result picks the branch (anxiety branch if anger)
    e_apathy = run("eac_esc_apathy", everyone).astype(bool)
    frustration = run("eac_esc_frustration", ~anger).astype(bool)
    e_covert = run("eac_esc_covert", ~anger & ~frustration).astype(bool)
    esc = []
    for a, fr, c, ap in zip(anger, frustration, e_covert, e_apathy):
        if a:
            esc.append(NAMES_BY_FLAGS[(0, 0, int(ap))])
        elif fr:
            esc.append("Frustration")
        else:
            esc.append(NAMES_BY_FLAGS[(1, int(c), int(ap))])
    return pd.Series(protect, index=idx), pd.Series(esc, index=idx)


if __name__ == "__main__":
    main()
