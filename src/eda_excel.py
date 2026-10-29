"""Styled Excel workbook for the EAC EDA (tables, bar charts, pie charts)."""

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, BLUE, MID, LIGHT = "1F3864", "2F5597", "8FAADC", "DAE3F3"
GREY, WHITE, GREEN, AMBER, RED = "F2F2F2", "FFFFFF", "E2EFDA", "FFF2CC", "F8CBAD"
PALETTE = [NAVY, BLUE, MID, "548235", "A9D18E", "BF9000", "C55A11", "7F7F7F"]
FONT = "Calibri"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ROLE_NAMES = {
    "feature": "Question (core)",
    "protected": "Question (protected attribute)",
    "other_label": "Other model output (final_*)",
    "identifier": "Identifier",
    "target": "Target",
}
# Is the role used as a model input in the baseline?
ROLE_INPUT = {
    "feature": "Yes",
    "protected": "No (removed)",
    "other_label": "Yes (kept)",
    "identifier": "No (removed)",
    "target": "No (this is what we predict)",
}
GUIDE = [
    ("Data type", "yes_no", "Only Yes or No answers."),
    ("Data type", "category", "Pick ONE option from a short fixed list (30 options or fewer), e.g. Rural / Suburban."),
    ("Data type", "multiselect", "Pick SEVERAL options; stored in one cell as 'Mom | Dad | Grandmother'. Becomes one 0/1 column per option."),
    ("Data type", "numeric", "A number, e.g. age or a 0-10 rating."),
    ("Data type", "short_text", "Typed-in SHORT answer (a word or a few words, median under 40 characters) with many different values, e.g. career or '3 adjectives'. Not a fixed list."),
    ("Data type", "free_text", "Typed-in LONG answer (sentences or paragraphs, median 40+ characters) where almost every answer is different, e.g. origin_freeform."),
    ("Data type", "all_blank", "Nobody answered this column."),
    ("Role", "Question (core)", "A normal questionnaire answer. Used as a model input."),
    ("Role", "Question (protected attribute)", "Age, gender, sexual orientation, religion, children's age/gender. REMOVED from inputs."),
    ("Role", "Other model output (final_*)", "Outputs of other models. final_osb and final_asb are real pipeline inputs. KEPT."),
    ("Role", "Identifier", "Respondent id, link, campaign. REMOVED."),
    ("Role", "Target", "final_eac_protect and final_eac_esc: what we predict."),
    ("Blank %", "blank %", "Share of respondents who left the column empty (100 minus coverage). Often 'not asked' because of skip logic."),
    ("Protect vs Escalating table", "rows / columns", "Each row is a Protect label, each column an Escalating label. Each cell counts respondents who have that exact pair. Darker blue = more respondents."),
]


def _safe(v):
    """Make a value safe for openpyxl."""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def _put(ws, row, col, value, *, bold=False, color="000000", size=11, fill=None, align=None, fmt=None, border=True):
    c = ws.cell(row, col)
    c.value = _safe(value)
    if isinstance(c.value, str) and c.value[:1] in "=+-@":
        c.data_type = "s"
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    return c


def title(ws, text, sub="", width=10):
    """Sheet banner."""
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    _put(ws, 1, 1, text, bold=True, color=WHITE, size=18, fill=NAVY, border=False)
    ws.row_dimensions[1].height = 34
    ws.cell(1, 1).alignment = Alignment(vertical="center", indent=1)
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        _put(ws, 2, 1, sub, color="595959", size=10, fill=LIGHT, border=False)
        ws.cell(2, 1).alignment = Alignment(vertical="center", indent=1)


def heading(ws, row, text, col=1):
    """Section heading."""
    c = _put(ws, row, col, text, bold=True, color=NAVY, size=13, border=False)
    c.border = Border(bottom=Side(style="medium", color=BLUE))


def table(ws, df, row, col=1, *, pct_cols=(), num_cols=(), center_cols=(), zebra=True):
    """Write a styled table; return (first_data_row, last_row)."""
    for j, name in enumerate(df.columns):
        _put(ws, row, col + j, str(name).replace("_", " ").title() if str(name).islower() else name,
             bold=True, color=WHITE, fill=BLUE, align="center")  # fmt: skip
    ws.row_dimensions[row].height = 22
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        fill = GREY if zebra and i % 2 == 0 else WHITE
        for j, v in enumerate(rec):
            name = df.columns[j]
            align = "center" if name in center_cols or name in pct_cols or name in num_cols else None
            fmt = "0.0" if name in pct_cols else ("#,##0" if name in num_cols else None)
            _put(ws, row + i, col + j, v, fill=fill, align=align, fmt=fmt)
    return row + 1, row + len(df)


def widths(ws, spec):
    """Set column widths from {letter: width}."""
    for letter, w in spec.items():
        ws.column_dimensions[letter].width = w


def bar(ws, anchor, ttl, cats, data, *, horizontal=False, color=BLUE, w=15, h=7.5, labels=True, stacked=False, colors=None):
    """Add a bar chart."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = ttl
    ch.style = 10
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.width, ch.height = w, h
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    else:
        ch.legend = None if len(ch.series) == 1 else ch.legend
    for k, s in enumerate(ch.series):
        s.graphicalProperties.solidFill = (colors or [color])[k % len(colors or [color])]
        s.graphicalProperties.line.solidFill = (colors or [color])[k % len(colors or [color])]
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
    ws.add_chart(ch, anchor)


def pie(ws, anchor, ttl, cats, data, *, w=13, h=7.5):
    """Add a pie chart with percentage labels."""
    ch = PieChart()
    ch.title = ttl
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.width, ch.height = w, h
    for k in range(cats.max_row - cats.min_row + 1):
        pt = DataPoint(idx=k)
        pt.graphicalProperties.solidFill = PALETTE[k % len(PALETTE)]
        ch.series[0].dPt.append(pt)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showPercent = True
    ch.dataLabels.showVal = False
    ch.dataLabels.showCatName = False
    ws.add_chart(ch, anchor)


def _action_fill(text: str) -> str:
    t = str(text)
    if t.startswith("keep"):
        return GREEN
    if t.startswith("drop"):
        return RED
    if t.startswith("skip") or t.startswith("review"):
        return AMBER
    return GREY


def _overview(wb, fname, summ, role_df, type_df, band_df, sugg_df, quality):
    ws = wb.active
    ws.title = "Overview"
    title(ws, "EAC Baseline: Data Overview", f"Source: {fname}  |  blank % = 100 - coverage (share of respondents who answered)", 14)
    n_rows = quality["Respondents analysed"]
    n_excl = quality["Rows excluded (branch rule breaks)"]
    kept = int(summ["suggestion"].str.startswith("keep").sum())
    dropped = int(summ["suggestion"].str.startswith("drop").sum())
    skipped = int(summ["suggestion"].str.startswith("skip").sum())
    heading(ws, 4, "Key figures")
    key = pd.DataFrame({
        "measure": ["Respondents analysed", "Rows excluded (branch rule breaks)", "Columns in file", "Kept for baseline", "Dropped", "Skipped (free text)"],
        "value": [n_rows, n_excl, len(summ), kept, dropped, skipped],
    })  # fmt: skip
    table(ws, key, 5, num_cols=("value",))

    role = role_df.copy()
    role["used_as_input"] = role["role"].map(ROLE_INPUT)
    role["role"] = role["role"].map(ROLE_NAMES).fillna(role["role"])
    r = 13
    blocks = [
        ("Columns by role (and whether the baseline uses them)", role, None),
        ("Question columns by data type", type_df, ("bar", "Question columns by data type")),
        ("Question columns by % blank", band_df, ("bar", "Question columns by % blank")),
        ("Baseline action per column", sugg_df, ("pie", "Baseline action")),
    ]
    for head, frame, chart in blocks:
        heading(ws, r, head)
        first, last = table(ws, frame, r + 1, num_cols=("columns",))
        if chart:
            cats = Reference(ws, min_col=1, min_row=first, max_row=last)
            data = Reference(ws, min_col=2, min_row=first - 1, max_row=last)
            if chart[0] == "bar":
                bar(ws, f"F{r}", chart[1], cats, data, w=16, h=7.5)
            else:
                pie(ws, f"F{r}", chart[1], cats, data, w=16, h=8)
        r += max(len(frame) + 3, 17)
    for col in "FGHIJKLMNO":
        ws.column_dimensions[col].width = 11
    heading(ws, r, "Data quality checks")
    qdf = pd.DataFrame({"check": list(quality.keys()), "result": [str(v) for v in quality.values()]})
    table(ws, qdf, r + 1)
    widths(ws, {"A": 52, "B": 16, "C": 38, "D": 3, "E": 3})


def _guide(wb):
    ws = wb.create_sheet("Guide")
    title(ws, "How to Read This Workbook", "Plain-language meaning of the types, roles and tables.", 3)
    df = pd.DataFrame(GUIDE, columns=["topic", "term", "meaning"])
    table(ws, df, 4)
    for row in ws.iter_rows(min_row=5, max_row=4 + len(df), min_col=3, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="center")
    widths(ws, {"A": 26, "B": 30, "C": 110})


def _columns(wb, summ):
    ws = wb.create_sheet("Columns")
    title(ws, "Every Column: Type, Blank %, Baseline Action", "Filter on any header. Green = keep, amber = skip or review, red = drop.", 11)
    df = summ.copy()
    df["group"] = df["group"].map(ROLE_NAMES).fillna(df["group"])
    df = df.rename(columns={"no": "#", "pct_blank": "blank_pct", "pct_answered": "answered_pct"})
    df = df[["#", "column", "group", "type", "n_answered", "n_blank", "blank_pct", "answered_pct", "n_unique", "suggestion"]]
    first, last = table(ws, df, 4, pct_cols=("blank_pct", "answered_pct"), num_cols=("n_answered", "n_blank", "n_unique"), center_cols=("#", "type"))
    sc = df.columns.get_loc("suggestion") + 1
    for rr in range(first, last + 1):
        ws.cell(rr, sc).fill = PatternFill("solid", fgColor=_action_fill(ws.cell(rr, sc).value))
    bc = get_column_letter(df.columns.get_loc("blank_pct") + 1)
    ws.conditional_formatting.add(f"{bc}{first}:{bc}{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=MID))
    ws.auto_filter.ref = f"A4:{get_column_letter(len(df.columns))}{last}"
    ws.freeze_panes = "C5"
    widths(ws, {"A": 6, "B": 70, "C": 30, "D": 13, "E": 11, "F": 9, "G": 10, "H": 12, "I": 10, "J": 40})


def _categories(wb, summ, long_df):
    ws = wb.create_sheet("Category index")
    title(ws, "Category, Yes/No and Multi-select Columns: At a Glance", "One row per question. 'Values' = number of distinct answer options.", 8)
    top = long_df[long_df["value_no"] == 1][["column_no", "value", "pct"]].rename(columns={"value": "top_value", "pct": "top_value_pct"})
    nvals = long_df.groupby("column_no").size().rename("values").reset_index()
    idx = summ[summ["type"].isin(["category", "yes_no", "multiselect"])][["no", "column", "type", "group", "pct_blank"]]
    idx = idx.merge(nvals, left_on="no", right_on="column_no").merge(top, on="column_no")
    idx["group"] = idx["group"].map(ROLE_NAMES).fillna(idx["group"])
    idx = idx[["no", "column", "type", "group", "values", "pct_blank", "top_value", "top_value_pct"]].rename(columns={"no": "#", "pct_blank": "blank_pct"})
    first, last = table(ws, idx, 4, pct_cols=("blank_pct", "top_value_pct"), num_cols=("values",), center_cols=("#", "type"))
    for col in ("blank_pct", "top_value_pct"):
        L = get_column_letter(idx.columns.get_loc(col) + 1)
        ws.conditional_formatting.add(f"{L}{first}:{L}{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=MID))
    ws.auto_filter.ref = f"A4:H{last}"
    ws.freeze_panes = "C5"
    widths(ws, {"A": 6, "B": 70, "C": 13, "D": 30, "E": 9, "F": 10, "G": 42, "H": 14})

    ws2 = wb.create_sheet("Category values")
    title(ws2, "Numbered Values for Each Question", "Rows are shaded by question. Filter on 'Column' to read one question.", 7)
    lv = long_df.rename(columns={"column_no": "#", "value_no": "value_#", "pct": "pct_of_answered"})[["#", "column", "type", "value_#", "value", "count", "pct_of_answered"]]
    for j, name in enumerate(lv.columns):
        _put(ws2, 4, 1 + j, str(name).replace("_", " ").title() if name.islower() else name, bold=True, color=WHITE, fill=BLUE, align="center")
    shade, prev = False, None
    for i, rec in enumerate(lv.itertuples(index=False), start=5):
        if rec[0] != prev:
            shade, prev = not shade, rec[0]
        fill = LIGHT if shade else WHITE
        for j, v in enumerate(rec):
            name = lv.columns[j]
            fmt = "0.0" if name == "pct_of_answered" else ("#,##0" if name == "count" else None)
            _put(ws2, i, 1 + j, v, fill=fill, align="center" if name in ("#", "value_#", "type", "count", "pct_of_answered") else None, fmt=fmt)
    end = 4 + len(lv)
    ws2.conditional_formatting.add(f"G5:G{end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=MID))
    ws2.auto_filter.ref = f"A4:G{end}"
    ws2.freeze_panes = "B5"
    widths(ws2, {"A": 6, "B": 60, "C": 13, "D": 9, "E": 55, "F": 9, "G": 15})


def _targets(wb, label_tables, ctab, nodes_df):
    ws = wb.create_sheet("Targets")
    title(ws, "EAC Labels and Training Rows per Model", "Labels cleaned to the agreed spelling (Covert Anger-Apathy).", 16)
    r = 4
    for name, tbl in label_tables.items():
        heading(ws, r, f"{name}: rows per label")
        first, last = table(ws, tbl, r + 1, num_cols=("rows",), pct_cols=("pct",), center_cols=("in_allowed_set",))
        pie(ws, f"G{r}", name, Reference(ws, min_col=1, min_row=first, max_row=last), Reference(ws, min_col=2, min_row=first - 1, max_row=last), w=13, h=7)
        r += 17

    heading(ws, r, "Protect label (rows) vs Escalating label (columns): number of respondents")
    ct = ctab.copy()
    ct.insert(0, "Protect \\ Escalating", ct.index)
    ct = ct.reset_index(drop=True)
    first, last = table(ws, ct, r + 1, num_cols=tuple(ct.columns[1:]), zebra=False)
    rng = f"B{first}:{get_column_letter(len(ct.columns))}{last}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=BLUE))
    r = last + 3

    heading(ws, r, "Training rows per cascade model")
    nd = nodes_df.copy()
    first, last = table(ws, nd, r + 1, num_cols=("eligible_rows", "positives", "negatives"), pct_cols=("pct_positive",))
    bar(ws, f"H{r}", "Positives vs negatives per model", Reference(ws, min_col=1, min_row=first, max_row=last),
        Reference(ws, min_col=3, max_col=4, min_row=first - 1, max_row=last), horizontal=True, stacked=True, w=18, h=8,
        colors=[NAVY, MID], labels=True)  # fmt: skip
    widths(ws, {"A": 26, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})
    for col in "GHIJKLMNOP":
        ws.column_dimensions[col].width = 12


def _scores(wb, scores_df, scores_esc_df, note):
    ws = wb.create_sheet("Covert scores")
    title(ws, "Covert Scores as the Pipeline Computes Them", note[:250], 10)
    if scores_df is None:
        return
    r = 4
    for label, frame in (("Protect", scores_df), ("Escalating", scores_esc_df)):
        heading(ws, r, f"Average score by {label} label")
        first, last = table(ws, frame, r + 1, num_cols=("rows",), pct_cols=("mean_overt_score", "mean_covert_score", "mean_bias_ratio"))
        bar(ws, f"G{r}", f"Average overt vs covert score by {label} label",
            Reference(ws, min_col=1, min_row=first, max_row=last),
            Reference(ws, min_col=3, max_col=4, min_row=first - 1, max_row=last), w=18, h=8, colors=[NAVY, "C55A11"], labels=False)  # fmt: skip
        r += 19
    widths(ws, {"A": 24, "B": 9, "C": 18, "D": 18, "E": 16})


def _breaks(wb, breaks_df):
    ws = wb.create_sheet("Branch rule breaks")
    title(ws, "Rows Excluded: Protect / Escalating Branch Rule Breaks", "Excluded from every statistic in this workbook; the source file is unchanged. Rule: Protect anger-family label goes with an Escalating Anxiety label, otherwise with an Anger-branch label. 'Excel row' = row in the source sheet.", 5)
    table(ws, breaks_df, 4, center_cols=("excel_row",))
    widths(ws, {"A": 12, "B": 26, "C": 26, "D": 40, "E": 40})


def _scatter(wb, summ, long_df, work_p, work_e, protect_order, esc_order):
    """Scatter charts: respondent scores by label, and question blank % vs top-answer share."""
    ws = wb.create_sheet("Scatter")
    wd = wb.create_sheet("Scatter data")
    title(ws, "Scatter Plots", "Each dot is one respondent (charts 1-2) or one question (chart 3). Respondent dots are nudged slightly so overlapping points stay visible.", 14)
    title(wd, "Data Behind the Scatter Plots", "Scores are jittered by up to 0.3 for display only.", 14)
    rng = np.random.default_rng(0)
    colors = [NAVY, "C55A11", "548235", "BF9000", "7030A0", "00B0F0", "7F7F7F"]

    def block(frame, label_col, order, col0, xcol, ycol, name_hdr):
        """Write a label-sorted block; return {label: (first_row, last_row)}."""
        frame = frame.copy()
        frame["_o"] = frame[label_col].map({lab: i for i, lab in enumerate(order)})
        frame = frame.sort_values(["_o"]).drop(columns="_o").reset_index(drop=True)
        for j, h in enumerate([name_hdr, label_col, xcol, ycol]):
            _put(wd, 4, col0 + j, h, bold=True, color=WHITE, fill=BLUE, align="center")
        ranges = {}
        for i, rec in enumerate(frame.itertuples(index=False), start=5):
            for j, v in enumerate(rec[:4]):
                wd.cell(i, col0 + j, _safe(v))
            lab = rec[1]
            ranges.setdefault(lab, [i, i])[1] = i
        return ranges

    def chart(anchor, ttl, xt, yt, ranges, order, xcol, ycol, xmax=10, ymax=10):
        ch = ScatterChart()
        ch.title, ch.style, ch.scatterStyle = ttl, 13, "lineMarker"
        ch.width, ch.height = 20, 11
        ch.x_axis.title, ch.y_axis.title = xt, yt
        ch.x_axis.scaling.min = ch.y_axis.scaling.min = 0
        ch.x_axis.scaling.max, ch.y_axis.scaling.max = xmax, ymax
        ch.x_axis.delete = ch.y_axis.delete = False
        for k, lab in enumerate(order):
            if lab not in ranges:
                continue
            a, b = ranges[lab]
            s = Series(Reference(wd, min_col=ycol, min_row=a, max_row=b), Reference(wd, min_col=xcol, min_row=a, max_row=b), title=lab)
            c = colors[k % len(colors)]
            s.marker = Marker(symbol="circle", size=5)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=c)
            s.marker.graphicalProperties.line.solidFill = c
            s.graphicalProperties.line.noFill = True
            ch.series.append(s)
        ws.add_chart(ch, anchor)

    if work_p is not None:
        for w, order, col0, label_col, anchor, which in (
            (work_p, protect_order, 1, "protect_label", "A5", "Protect"),
            (work_e, esc_order, 6, "escalating_label", "A29", "Escalating"),
        ):
            f = pd.DataFrame({
                "excel_row": w.index + 2, label_col: w["label"].values,
                "overt_jittered": w["overt_score"].values + rng.uniform(-0.3, 0.3, len(w)),
                "covert_jittered": w["covert_score"].values + rng.uniform(-0.3, 0.3, len(w)),
            })  # fmt: skip
            rngs = block(f, label_col, order, col0, "overt_jittered", "covert_jittered", "excel_row")
            chart(anchor, f"Overt vs covert score by {which} label (each dot = one respondent)",
                  "Overt score (0-9)", "Covert score (0-9)", rngs, order, col0 + 2, col0 + 3)  # fmt: skip

    # chart 3: question usefulness
    top = long_df[long_df["value_no"] == 1][["column_no", "pct"]].rename(columns={"pct": "top_answer_pct"})
    q = summ[summ["type"].isin(["yes_no", "category"])][["no", "type", "pct_blank"]].merge(top, left_on="no", right_on="column_no")
    q = q[["no", "type", "pct_blank", "top_answer_pct"]]
    q_order = ["yes_no", "category"]
    qr = block(q, "type", q_order, 11, "pct_blank", "top_answer_pct", "question_no")
    chart("A53", "Yes/No and category questions: % blank vs share of the most common answer",
          "% blank", "Most common answer, % of those who answered", qr, q_order, 13, 14, xmax=100, ymax=100)
    near_const = int((q["top_answer_pct"] >= 95).sum())
    ws.merge_cells("A77:N79")
    _put(ws, 77, 1, f"How to read chart 3: dots near the top (95%+) are questions where almost everyone gives the same answer, so they carry little signal ({near_const} questions). Dots to the right are mostly blank questions. Find a dot's question by its number in the 'Category index' sheet.", color="404040", border=False)
    ws.cell(77, 1).alignment = Alignment(wrap_text=True, vertical="top")
    widths(wd, {"A": 11, "B": 22, "C": 15, "D": 15, "F": 11, "G": 22, "H": 15, "I": 15, "K": 12, "L": 12, "M": 12, "N": 16})


def build_workbook(path, fname, summ, role_df, type_df, band_df, sugg_df, long_df, label_tables, ctab, nodes_df, scores_df, scores_esc_df, scores_note, quality, breaks_df, work_p=None, work_e=None):
    """Create the styled EDA workbook at ``path``."""
    wb = Workbook()
    _overview(wb, fname, summ, role_df, type_df, band_df, sugg_df, quality)
    _guide(wb)
    _targets(wb, label_tables, ctab, nodes_df)
    _scores(wb, scores_df, scores_esc_df, scores_note)
    _scatter(wb, summ, long_df, work_p, work_e, list(scores_df["label"]) if scores_df is not None else [], list(scores_esc_df["label"]) if scores_esc_df is not None else [])
    _breaks(wb, breaks_df)
    _columns(wb, summ)
    _categories(wb, summ, long_df)
    wb.save(path)
